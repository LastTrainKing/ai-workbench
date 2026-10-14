"""Demo 2 — the other two formats, through the same closed loop.

Demo 1 proved the loop with a presentation. The product promises Word and Excel
too, and "it also supports docx" is a claim that has to survive being checked the
same way: produce the file, then open it with an independent library and count what
is inside.

    Word   .docx   > 30 KB, ≥ 5 pages, a real heading hierarchy
    Excel  .xlsx   a formula that openpyxl reads back starting with "=",
                   and at least one chart

Page count for the Word file is a *real* page count, not an estimate: LibreOffice
converts it to PDF and the PDF is counted. A .docx does not store its own page
count, and reporting "roughly five pages" would be the kind of number this project
has spent four rounds learning not to print.

Run:  python scripts/demo_2_word_excel.py
"""

from __future__ import annotations

import itertools
import json
import shutil
import subprocess
import tempfile
import time
from pathlib import Path
from typing import Any

import httpx

ROOT = Path(__file__).resolve().parents[1]
BASE = "http://127.0.0.1:8000/api/v1"
POLL_SECONDS = 8
MAX_POLLS = 120
TERMINAL = {"COMPLETED", "PARTIAL", "FAILED", "CANCELLED", "LIMIT_REACHED", "WAITING_USER"}

WORD_GOAL = (
    "用我导入的毕业设计资料，写一份《系统设计说明书》，必须用 document_write 生成 .docx 文件。"
    "要求：至少 8 个章节，每章用 heading 做标题，正文内容充实具体，总字数不少于 4000 字。"
    "最后一步要回读该文件核对段落数。"
)
EXCEL_GOAL = (
    "把 AI Workbench 的工具清单整理成一份 Excel：每行一个模块（模块名、工具数量、权限级别），"
    "至少 8 行数据；最后一行用真实公式求和工具总数（formulas 里写 =SUM(...)）；"
    "并插入一个柱状图。必须用 document_write 生成 .xlsx 文件，然后回读核对。"
)

SOFFICE_CANDIDATES = (
    Path(r"C:\Program Files\LibreOffice\program\soffice.exe"),
    Path(r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"),
)


def line(title: str) -> None:
    print(f"\n{'=' * 78}\n  {title}\n{'=' * 78}")


def data(response: httpx.Response) -> Any:
    payload = response.json()
    if not payload.get("success", True):
        raise RuntimeError(f"HTTP {response.status_code}: {json.dumps(payload, ensure_ascii=False)[:300]}")
    return payload["data"]


def login(client: httpx.Client) -> dict[str, str]:
    token = client.post(
        f"{BASE}/auth/login", json={"identifier": "demo", "password": "demo1234"}
    ).json()["data"]["access_token"]
    return {"Authorization": f"Bearer {token}"}


def run_build(
    client: httpx.Client, headers: dict[str, str], *, goal: str, kind: str, title: str
) -> dict[str, Any]:
    created = data(
        client.post(
            f"{BASE}/build-runs",
            headers=headers,
            json={"goal": goal, "deliverable_kind": kind, "title": title},
        )
    )
    run_id = created["id"]
    print(f"  BuildRun {run_id[:8]}  「{title}」")
    started = client.post(f"{BASE}/build-runs/{run_id}/execute", headers=headers)
    print(f"  execute → HTTP {started.status_code}")
    if started.status_code != 200:
        print(f"  {started.text[:300]}")
        return {"id": run_id, "status": "EXECUTE_FAILED", "deliverables": []}

    for index in range(MAX_POLLS):
        time.sleep(POLL_SECONDS)
        run = data(client.get(f"{BASE}/build-runs/{run_id}", headers=headers))
        if index % 4 == 0:
            print(f"    … {run['status']}  {run['steps_done']}/{run['steps_total']}")
        if run["status"] in TERMINAL:
            print(f"  终态 {run['status']}  步骤 {run['steps_done']}/{run['steps_total']}")
            break
    run["deliverables"] = data(
        client.get(f"{BASE}/deliverables", headers=headers, params={"build_run_id": run_id})
    )["deliverables"]
    return run


def soffice() -> Path | None:
    for candidate in SOFFICE_CANDIDATES:
        if candidate.is_file():
            return candidate
    found = shutil.which("soffice")
    return Path(found) if found else None


def docx_page_count(path: Path) -> tuple[int | None, str]:
    """A real page count, by converting to PDF and counting it."""
    exe = soffice()
    if exe is None:
        return None, "LibreOffice 未安装，无法得到真实页数"
    with tempfile.TemporaryDirectory() as tmp:
        result = subprocess.run(
            [str(exe), "--headless", "--convert-to", "pdf", "--outdir", tmp, str(path)],
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            timeout=180,
            check=False,
        )
        pdf = Path(tmp) / (path.stem + ".pdf")
        if not pdf.is_file():
            return None, f"转换失败：{(result.stdout or '')[:160]} {(result.stderr or '')[:160]}"
        import pypdf

        return len(pypdf.PdfReader(str(pdf)).pages), ""


def inspect_docx(path: Path) -> dict[str, Any]:
    from docx import Document

    document = Document(str(path))
    paragraphs = [p for p in document.paragraphs if p.text.strip()]
    headings: list[tuple[int, str]] = []
    for paragraph in paragraphs:
        name = paragraph.style.name or ""
        if name.startswith("Heading"):
            try:
                level = int(name.split()[-1])
            except ValueError:
                level = 1
            headings.append((level, paragraph.text.strip()))
    chars = sum(len(p.text) for p in paragraphs)
    pages, note = docx_page_count(path)

    # A hierarchy is only a hierarchy if it nests: at least two levels, and the
    # first heading is not a sub-heading of something that was never introduced.
    levels = {level for level, _ in headings}
    hierarchy_ok = (
        len(headings) >= 5
        and len(levels) >= 2
        and headings[0][0] == min(levels)
        and all(
            level <= previous + 1
            for (previous, _), (level, _) in itertools.pairwise(headings)
        )
    )
    return {
        "path": str(path),
        "size_bytes": path.stat().st_size,
        "paragraphs": len(paragraphs),
        "chars": chars,
        "headings": len(headings),
        "heading_levels": sorted(levels),
        "heading_sample": [f"H{level} {text[:34]}" for level, text in headings[:6]],
        "pages": pages,
        "page_note": note,
        "hierarchy_ok": hierarchy_ok,
    }


def inspect_xlsx(path: Path) -> dict[str, Any]:
    from openpyxl import load_workbook

    workbook = load_workbook(str(path), data_only=False)
    formulas: list[tuple[str, str]] = []
    rows = 0
    sheets: list[str] = []
    for name in workbook.sheetnames:
        sheet = workbook[name]
        sheets.append(f"{name}({sheet.max_row}×{sheet.max_column})")
        rows += sheet.max_row
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    formulas.append((f"{name}!{cell.coordinate}", cell.value))
    charts = sum(len(workbook[n]._charts) for n in workbook.sheetnames)
    return {
        "path": str(path),
        "size_bytes": path.stat().st_size,
        "sheets": sheets,
        "rows": rows,
        "formulas": formulas,
        "charts": charts,
    }


def main() -> int:
    client = httpx.Client(timeout=180)
    headers = login(client)
    results: dict[str, Any] = {}

    # ------------------------------------------------------------------ word --
    line("Demo 2a —— Word：真实 .docx，> 30 KB / ≥ 5 页 / 标题层级正确")
    word_run = run_build(client, headers, goal=WORD_GOAL, kind="docx", title="系统设计说明书（Demo 2）")
    word_item = next((d for d in word_run.get("deliverables") or [] if d["kind"] == "docx"), None)
    if word_item is None:
        print("  没有登记 .docx 产出物 —— 如实记录为 FAILED")
        results["word"] = {"run": word_run, "passed": False, "reason": "no docx deliverable"}
    else:
        word_path = Path(word_item["path"])
        print(f"  产出物 {word_path.name}  {word_item['size_bytes']:,} B  verified={word_item['verified']}")
        facts = inspect_docx(word_path)
        print(f"  段落 {facts['paragraphs']}  字数 {facts['chars']}  标题 {facts['headings']} 个")
        print(f"  标题层级：{facts['heading_levels']}")
        for sample in facts["heading_sample"]:
            print(f"    {sample}")
        print(f"  真实页数（LibreOffice → PDF）：{facts['pages']} {facts['page_note']}")
        checks = {
            "产出 .docx": word_path.suffix == ".docx" and word_path.is_file(),
            "大小 > 30 KB": facts["size_bytes"] > 30_000,
            "页数 ≥ 5（真实转换页数）": bool(facts["pages"] and facts["pages"] >= 5),
            "标题层级正确（≥2 级且能嵌套）": facts["hierarchy_ok"],
            "运行终态 COMPLETED": word_run["status"] == "COMPLETED",
            "产出物机器校验通过": bool(word_item["verified"]),
        }
        for name, ok in checks.items():
            print(f"    [{'OK' if ok else 'X'}] {name}")
        results["word"] = {
            "run_status": word_run["status"],
            "deliverable": word_item,
            "facts": facts,
            "checks": checks,
            "passed": all(checks.values()),
        }

    # ----------------------------------------------------------------- excel --
    line("Demo 2b —— Excel：真实公式 + 至少 1 张图表")
    excel_run = run_build(client, headers, goal=EXCEL_GOAL, kind="xlsx", title="工具清单（Demo 2）")
    excel_item = next((d for d in excel_run.get("deliverables") or [] if d["kind"] == "xlsx"), None)
    if excel_item is None:
        print("  没有登记 .xlsx 产出物 —— 如实记录为 FAILED")
        results["excel"] = {"run": excel_run, "passed": False, "reason": "no xlsx deliverable"}
    else:
        excel_path = Path(excel_item["path"])
        print(f"  产出物 {excel_path.name}  {excel_item['size_bytes']:,} B  verified={excel_item['verified']}")
        facts = inspect_xlsx(excel_path)
        print(f"  工作表：{facts['sheets']}  总行数 {facts['rows']}")
        print(f"  真实公式 {len(facts['formulas'])} 个：")
        for coordinate, formula in facts["formulas"][:5]:
            print(f"    {coordinate} = {formula}")
        print(f"  图表 {facts['charts']} 个")
        checks = {
            "产出 .xlsx": excel_path.suffix == ".xlsx" and excel_path.is_file(),
            "含真实公式（openpyxl 读到以 = 开头）": bool(facts["formulas"]),
            "至少 1 张图表": facts["charts"] >= 1,
            "有数据行（≥ 5 行）": facts["rows"] >= 5,
            "运行终态 COMPLETED": excel_run["status"] == "COMPLETED",
            "产出物机器校验通过": bool(excel_item["verified"]),
        }
        for name, ok in checks.items():
            print(f"    [{'OK' if ok else 'X'}] {name}")
        results["excel"] = {
            "run_status": excel_run["status"],
            "deliverable": excel_item,
            "facts": facts,
            "checks": checks,
            "passed": all(checks.values()),
        }

    line("结论")
    overall = all(part.get("passed") for part in results.values())
    for name, part in results.items():
        print(f"  [{'PASS' if part.get('passed') else 'FAILED'}] Demo 2{name[0]} —— {name}")
    out = ROOT / ".runtime" / "demo2-record.json"
    out.write_text(json.dumps(results, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    print(f"\n  完整记录：{out.relative_to(ROOT)}")
    return 0 if overall else 1


if __name__ == "__main__":
    raise SystemExit(main())
