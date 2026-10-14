"""Preference acceptance: the four criteria, run for real.

The claim being tested is not "the prompt mentions the precedence rule". It is that
the product, given a task and a stated preference, does the *right* thing and says
what it did:

    判据 1  15 分钟 + 页数更少   → 10~14 页，并说明这是硬约束下的下限
    判据 2  15 分钟 + 图表更多   → images > 0
    判据 3  3 分钟 + 内容详尽    → 服从时间约束，并说明为什么做不到详尽
    判据 4  模板不存在           → 明确报告做不到，不静默忽略

Every number below is read from the produced .pptx with python-pptx, and the
explanations are read from the run's own report. Nothing here is taken from what the
agent says about itself.

Run:  python scripts/verify_preferences.py
"""

from __future__ import annotations

import asyncio
import json
import shutil
import sys
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "backend"))

SCRATCH = ROOT / ".runtime" / "pref-check"

#: Marker so the temporary preference rows can be removed precisely.
TEMP_MARK = "[acceptance-check]"

TITLE = "毕业设计答辩"
GOAL_15 = (
    "用资料做一份 15 分钟的毕业设计答辩 PPT，包含系统架构、数据库设计、AI 模块三部分。"
)
GOAL_3 = "做一份 3 分钟的快速汇报 PPT，介绍这套系统是做什么的。"
GOAL_PLAIN = "做一份介绍本项目架构的 PPT。"


def line(title: str) -> None:
    print(f"\n{'=' * 78}\n  {title}\n{'=' * 78}")


def measure(path: Path) -> dict[str, Any]:
    """Read the deck back and count what is actually in it."""
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    deck = Presentation(str(path))
    slides = len(deck.slides)
    images = sum(
        1
        for slide in deck.slides
        for shape in slide.shapes
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
    )
    chars = sum(
        len(shape.text_frame.text)
        for slide in deck.slides
        for shape in slide.shapes
        if getattr(shape, "has_text_frame", False)
    )
    titles = []
    for slide in deck.slides:
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
                titles.append(shape.text_frame.text.strip().split("\n")[0][:40])
                break
    return {
        "path": str(path),
        "size_bytes": path.stat().st_size,
        "slides": slides,
        "images": images,
        "chars": chars,
        "titles": titles,
    }


async def with_temporary_preferences(texts: list[str]) -> None:
    """Insert preference memories, used by the criteria that need one added."""
    from app.infrastructure.database import session_scope
    from app.models.context import MemoryEntry

    async with session_scope() as session:
        from app.models.user import User
        from sqlalchemy import select

        user = (await session.execute(select(User).limit(1))).scalar_one()
        for text in texts:
            session.add(
                MemoryEntry(
                    user_id=user.id,
                    kind="preference",
                    scope="global",
                    content=f"{TEMP_MARK} {text}",
                    confidence=0.95,
                    source="acceptance",
                    status="active",
                    confirmed_by_user=True,
                )
            )
        await session.commit()


async def clear_temporary_preferences() -> int:
    from app.infrastructure.database import session_scope
    from app.models.context import MemoryEntry
    from sqlalchemy import delete

    async with session_scope() as session:
        result = await session.execute(
            delete(MemoryEntry).where(MemoryEntry.content.like(f"{TEMP_MARK}%"))
        )
        await session.commit()
        return result.rowcount or 0


async def run_case(label: str, goal: str) -> dict[str, Any]:
    """Drive the real agent for one criterion and return the record."""
    from app.ai.agent.autonomous import AutonomousAgent
    from app.ai.agent.control import RunController
    from app.ai.providers.factory import build_llm_provider
    from app.ai.tools.audit import make_audit_sink
    from app.ai.tools.confirmation import get_broker
    from app.ai.tools.executor import ResourceLimits, ToolExecutor
    from app.ai.tools.policy_store import build_policy, capability_flags
    from app.ai.tools.registry import build_computer_registry
    from app.ai.tools.unified import UnifiedToolRegistry
    from app.infrastructure.database import session_scope
    from app.models.user import User
    from sqlalchemy import select

    workdir = SCRATCH / label
    if workdir.exists():
        shutil.rmtree(workdir)
    workdir.mkdir(parents=True)

    async with session_scope() as session:
        owner = (await session.execute(select(User).limit(1))).scalar_one()

        policy = build_policy()
        policy.workspace = workdir
        policy.allowed_paths = [*policy.allowed_paths, workdir]

        business = __import__("app.ai.agent.tools", fromlist=["get_tool_registry"]).get_tool_registry()
        computer = build_computer_registry(policy, enabled=capability_flags())
        limits = ResourceLimits(
            max_output_bytes=policy.max_output_bytes,
            max_files_modified=policy.max_files_modified,
            max_tool_calls=policy.max_tool_calls,
            max_runtime_seconds=policy.max_runtime_seconds,
        )
        executor = ToolExecutor(policy, audit=make_audit_sink(session))
        registry = UnifiedToolRegistry(
            business, computer, executor, confirm=get_broker().request, limits=limits
        )

        llm, _info = build_llm_provider()
        controller = RunController(run_id=f"pref-{label}")
        agent = AutonomousAgent(
            llm=llm,
            registry=registry,
            policy=policy,
            controller=controller,
            session=session,
            owner_id=owner.id,
            limits=limits,
            working_directory=str(workdir),
            deliverable_kind="pptx",
        )

        print(f"  目标：{goal}")
        result = await agent.run(goal)
        print(f"  终态：{result.state.value}")
        # Printed *after* the run: the constraints are derived inside `run()`, so
        # reading them before it shows the empty defaults and reads like a bug in
        # the derivation rather than in this print statement.
        print(f"  解析出的硬约束：{agent.constraints.as_dict()}")

        print("  偏好处理结果：")
        for outcome in agent.preferences.outcomes:
            print(f"    [{outcome.verdict}] 「{outcome.text[:56]}」→ {outcome.detail[:110]}")
        print("  取舍说明（运行报告里的原文）：")
        for row in result.answer.split("\n"):
            if row.startswith("- ") or "取舍说明" in row:
                print(f"    {row[:150]}")

        return {
            "label": label,
            "goal": goal,
            "state": result.state.value,
            "answer": result.answer,
            "constraints": agent.constraints.as_dict(),
            "outcomes": [o.as_dict() for o in agent.preferences.outcomes],
            "preferences": agent.preferences.as_dict(),
            "workdir": str(workdir),
            "events": [
                {"kind": e.kind, "message": e.message} for e in controller.events
            ],
            "tool_calls": result.tool_calls,
        }


def find_deck(workdir: Path) -> Path | None:
    decks = sorted(workdir.rglob("*.pptx"), key=lambda p: p.stat().st_mtime, reverse=True)
    return decks[0] if decks else None


async def criterion_1() -> dict[str, Any]:
    line("判据 1 —— 「15 分钟答辩 PPT」+ 偏好「页数更少」→ 10~14 页 + 取舍说明")
    record = await run_case("c1", GOAL_15)
    deck = find_deck(Path(record["workdir"]))
    record["deck"] = measure(deck) if deck else None
    if record["deck"]:
        print(f"\n  产出：{record['deck']['slides']} 页 / {record['deck']['images']} 图 / "
              f"{record['deck']['chars']} 字 / {record['deck']['size_bytes']:,} B")
    report = record["answer"]
    checks = {
        "产出了 .pptx": record["deck"] is not None,
        "页数落在 10~14": bool(record["deck"] and 10 <= record["deck"]["slides"] <= 14),
        "终态 COMPLETED": record["state"] == "COMPLETED",
        "报告里有取舍说明": "取舍说明" in report,
        "说明这是 15 分钟约束下的最小值": "最小值" in report or "下限" in report,
    }
    for name, ok in checks.items():
        print(f"    [{'OK' if ok else 'X'}] {name}")
    record["checks"] = checks
    record["passed"] = all(checks.values())
    return record


async def criterion_2() -> dict[str, Any]:
    line("判据 2 —— 「15 分钟答辩 PPT」+ 偏好「图表更多」→ images > 0")
    record = await run_case("c2", GOAL_15)
    deck = find_deck(Path(record["workdir"]))
    record["deck"] = measure(deck) if deck else None
    if record["deck"]:
        print(f"\n  产出：{record['deck']['slides']} 页 / {record['deck']['images']} 图 / "
              f"{record['deck']['chars']} 字 / {record['deck']['size_bytes']:,} B")
    checks = {
        "产出了 .pptx": record["deck"] is not None,
        "images > 0": bool(record["deck"] and record["deck"]["images"] > 0),
        "图表偏好被判为 applied": any(
            o["verdict"] == "applied" and "图" in o["text"] for o in record["outcomes"]
        ),
        "报告里说明了图表偏好生效": "图表偏好生效" in record["answer"],
    }
    for name, ok in checks.items():
        print(f"    [{'OK' if ok else 'X'}] {name}")
    record["checks"] = checks
    record["passed"] = all(checks.values())
    return record


async def criterion_3() -> dict[str, Any]:
    line("判据 3 —— 「3 分钟快速汇报」+ 偏好「内容详尽」→ 服从时间约束 + 说明为什么没做详尽")
    await with_temporary_preferences(["内容要尽量详尽，不要遗漏任何细节"])
    try:
        record = await run_case("c3", GOAL_3)
    finally:
        removed = await clear_temporary_preferences()
        print(f"\n  （清理临时偏好 {removed} 条）")

    deck = find_deck(Path(record["workdir"]))
    record["deck"] = measure(deck) if deck else None
    if record["deck"]:
        print(f"\n  产出：{record['deck']['slides']} 页 / {record['deck']['images']} 图 / "
              f"{record['deck']['chars']} 字 / {record['deck']['size_bytes']:,} B")
    checks = {
        "产出了 .pptx": record["deck"] is not None,
        "页数落在 3~5（服从 3 分钟约束）": bool(
            record["deck"] and 3 <= record["deck"]["slides"] <= 5
        ),
        "详尽偏好被判为 bounded": any(
            o["verdict"] == "bounded" and "详尽" in o["text"] for o in record["outcomes"]
        ),
        "报告里说明了为什么没做详尽": "做不到详尽" in record["answer"],
    }
    for name, ok in checks.items():
        print(f"    [{'OK' if ok else 'X'}] {name}")
    record["checks"] = checks
    record["passed"] = all(checks.values())
    return record


async def criterion_4() -> dict[str, Any]:
    line("判据 4 —— 无法满足的偏好（用公司模板，但模板不存在）→ 明确报告")
    await with_temporary_preferences(["必须使用公司模板，套用公司统一的 PPT 模板"])
    try:
        record = await run_case("c4", GOAL_PLAIN)
    finally:
        removed = await clear_temporary_preferences()
        print(f"\n  （清理临时偏好 {removed} 条）")

    deck = find_deck(Path(record["workdir"]))
    record["deck"] = measure(deck) if deck else None
    events = [e["message"] for e in record["events"]]
    reported = [m for m in events if "偏好无法满足" in m]
    checks = {
        "模板偏好被判为 unsatisfiable": any(
            o["verdict"] == "unsatisfiable" and "模板" in o["text"] for o in record["outcomes"]
        ),
        "事件流里有明确的「偏好无法满足」报告": bool(reported),
        "报告里说明了做不到": "做不到" in record["answer"] or "无法满足" in record["answer"],
        "没有静默忽略（不是既没报告也没生效）": bool(reported),
    }
    for name, ok in checks.items():
        print(f"    [{'OK' if ok else 'X'}] {name}")
    print(f"\n  报告原文片段：{reported[0][:160] if reported else '（无）'}")
    record["checks"] = checks
    record["passed"] = all(checks.values())
    return record


async def main() -> int:
    line("偏好语义验收 —— 四个判据，真实运行")
    SCRATCH.mkdir(parents=True, exist_ok=True)

    results: list[dict[str, Any]] = []
    for runner in (criterion_1, criterion_2, criterion_3, criterion_4):
        results.append(await runner())

    line("结论")
    for record in results:
        print(f"  [{'PASS' if record['passed'] else 'FAILED'}] {record['label']}")

    (SCRATCH / "results.json").write_text(
        json.dumps(results, ensure_ascii=False, indent=2, default=str), encoding="utf-8"
    )
    print(f"\n  完整结果已写入 {SCRATCH / 'results.json'}")
    return 0 if all(record["passed"] for record in results) else 1


if __name__ == "__main__":
    raise SystemExit(asyncio.run(main()))
